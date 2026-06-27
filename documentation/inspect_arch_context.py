"""Look at the content around the architecture section."""
from docx import Document

doc = Document("documentation/MidSem_Report_2024AA05606_styled.docx")
for i in range(85, min(95, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f"[{i}]: {p.text[:80]}")
