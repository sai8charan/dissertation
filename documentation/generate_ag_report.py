"""
documentation/generate_ag_report.py
-------------------------------------
Apply all faculty-requested updates (same as update_midsem_report.py) to
MidSem_Report_2024AA05606.docx and save the result as
MidSem_Report_2024AA05606_ag.docx.

Run from dissertation root:
    python documentation/generate_ag_report.py
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT     = Path(__file__).parent.parent
SRC_PATH = ROOT / "documentation" / "MidSem_Report_2024AA05606.docx"
OUT_PATH = ROOT / "documentation" / "MidSem_Report_2024AA05606_ag.docx"
CH2_MD   = Path(__file__).parent / "chapters" / "chapter2_literature_review.md"

# ── Content constants (from guidelines / CHANGELOG) ───────────────────────────

CONTRIBUTIONS = (
    "The specific contributions of this dissertation are: (1) a hybrid BM25 + "
    "sentence-embedding evidence retrieval module that selects salient context "
    "within BART's input budget; (2) a Best-of-N generation strategy combined "
    "with NLI-based sentence-level factual verification producing a Factual "
    "Consistency Score; (3) a preference reranker with extractive fallback that "
    "guarantees the system never silently returns a low-confidence summary; and "
    "(4) a systematic ablation study and human evaluation validating each "
    "component's contribution."
)

OBJECTIVES_TOP = [
    "Achieve target ROUGE-2 >= 18.0 on the CNN/DailyMail test set (above the Lead-3 benchmark)",
    "Achieve target NLI-FCS >= 0.65 for the proposed pipeline versus <= 0.55 for single-stage BART fine-tuned",
    "Maintain extractive fallback rate below 15% of test documents",
    "Achieve target human evaluation mean factual accuracy >= 4.0/5.0 with Cohen's Kappa >= 0.6",
    "Survey extractive and abstractive methods and document findings (standalone literature review in Chapter 2)",
    "Conduct a multilingual feasibility study on Hindi XL-Sum (translate-then-summarise vs direct mBART)",
    "Document all findings, limitations, and future directions in the final dissertation",
]

OBJECTIVES_CH1 = [
    "Survey and compare extractive and abstractive summarisation methods (Chapter 2); establish Lead-3 and zero-shot BART baselines",
    "Design and implement a two-stage pipeline: hybrid BM25 + embedding retrieval with fine-tuned BART Best-of-N generation",
    "Add NLI-based verification, preference reranking, and extractive fallback (target NLI-FCS >= 0.65; fallback rate < 15%)",
    "Evaluate using ROUGE, BERTScore, and NLI-FCS with ablation studies (target ROUGE-2 >= 18.0)",
    "Conduct human evaluation on 50 samples, 3 raters (target factual accuracy >= 4.0/5.0, Cohen's Kappa >= 0.6)",
    "Conduct multilingual feasibility study and document all findings in the final dissertation",
]

TARGETS_NOTE = (
    "Upon completion of experiments, the word 'target' and the numeric thresholds "
    "in this section will be replaced with achieved results in Chapter 5."
)

LITERATURE_POINTER = (
    "A standalone literature review covering extractive and abstractive summarisation, "
    "pretrained transformer models (BART, PEGASUS, mBART), factuality evaluation "
    "(ROUGE, BERTScore, SummaC, ACL 2026 stress-testing), and retrieval-augmented "
    "generation is presented in Chapter 2. Recent 2026 studies (Liu et al., 2026; "
    "Mujahid et al., 2026) confirm that summarisation remains an active research "
    "problem on factual consistency, motivating the verification stage of this pipeline."
)

CH2_MARKERS = (
    "Chapter 2: Literature Review",
    "2.1 Extractive Summarisation",
    "2.6 Research Gap and Contribution",
    "This dissertation fills the gap by chaining four modular stages",
)

# ── Helpers (identical to update_midsem_report.py) ────────────────────────────

def _delete_paragraph(para: Paragraph) -> None:
    element = para._element
    element.getparent().remove(element)


def _insert_paragraph_before(ref: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    ref._p.addprevious(new_p)
    para = Paragraph(new_p, ref._parent)
    if text:
        para.add_run(text)
    if style:
        try:
            para.style = style
        except KeyError:
            pass
    return para


def _set_paragraph_text(para: Paragraph, text: str, italic: bool = False) -> None:
    para.clear()
    if text:
        run = para.add_run(text)
        run.italic = italic


def _find_paragraph(doc: Document, prefix: str) -> int:
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip().startswith(prefix):
            return i
    raise ValueError(f"Paragraph starting with {prefix!r} not found")


def _replace_list_items(doc: Document, start_prefix: str, stop_prefixes: tuple, items: list[str]) -> None:
    start = _find_paragraph(doc, start_prefix)
    stop = len(doc.paragraphs)
    for prefix in stop_prefixes:
        try:
            stop = min(stop, _find_paragraph(doc, prefix))
        except ValueError:
            pass

    list_indices = []
    for i in range(start + 1, stop):
        para = doc.paragraphs[i]
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else ""
        if "Heading" in style_name:
            break
        list_indices.append(i)

    for idx, text in zip(list_indices, items):
        _set_paragraph_text(doc.paragraphs[idx], text)
    for idx in list_indices[len(items):]:
        _set_paragraph_text(doc.paragraphs[idx], "")


def _renumber_chapters(doc: Document) -> None:
    replacements = [
        ("Chapter 6: Bibliography / References", "Chapter 7: Bibliography / References"),
        ("Chapter 6: Bibliography", "Chapter 7: Bibliography / References"),
        ("Chapter 5: Directions for Future Work", "Chapter 6: Directions for Future Work"),
        ("Chapter 4: Preliminary Results and Objectives Achieved till Midterm",
         "Chapter 5: Preliminary Results and Objectives Achieved till Midterm"),
        ("Chapter 4: Preliminary Results", "Chapter 5: Preliminary Results and Objectives Achieved till Midterm"),
        ("Chapter 3: Methodology and Pipeline Implementation",
         "Chapter 4: Methodology and Pipeline Implementation"),
        ("Chapter 2: System Architecture and Tools Used",
         "Chapter 3: System Architecture and Tools Used"),
    ]
    for old, new in replacements:
        for para in doc.paragraphs:
            if para.text.strip() == old:
                _set_paragraph_text(para, new)

    section_ranges = [
        ("Chapter 3: System Architecture", "Chapter 4: Methodology", "3"),
        ("Chapter 4: Methodology", "Chapter 5: Preliminary", "4"),
        ("Chapter 5: Preliminary", "Chapter 6: Directions", "5"),
    ]
    for start_prefix, end_prefix, num in section_ranges:
        try:
            start = _find_paragraph(doc, start_prefix)
            end = _find_paragraph(doc, end_prefix)
        except ValueError:
            continue
        for i in range(start + 1, end):
            text = doc.paragraphs[i].text.strip()
            if re.match(r"^\d+\.\d+", text):
                _set_paragraph_text(
                    doc.paragraphs[i],
                    re.sub(r"^\d+\.", f"{num}.", text, count=1),
                )


def _parse_chapter2_blocks() -> list:
    lines = CH2_MD.read_text(encoding="utf-8").splitlines()
    blocks = []
    for line in lines:
        if not line.strip():
            continue
        if line.startswith("# "):
            blocks.append(("heading1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("heading2", line[3:].strip()))
        elif line.startswith("|") and "---" not in line:
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(line.strip())
            else:
                blocks.append(("table", [line.strip()]))
        else:
            blocks.append(("body", line.strip()))
    return blocks


def _insert_table_before(ref: Paragraph, row_lines: list[str]) -> None:
    rows = [[c.strip() for c in r.split("|") if c.strip()] for r in row_lines]
    if not rows:
        return
    n_cols = max(len(r) for r in rows)

    tbl = OxmlElement("w:tbl")
    for row in rows:
        tr = OxmlElement("w:tr")
        for col_idx in range(n_cols):
            tc = OxmlElement("w:tc")
            p  = OxmlElement("w:p")
            r  = OxmlElement("w:r")
            t  = OxmlElement("w:t")
            t.text = row[col_idx] if col_idx < len(row) else ""
            r.append(t)
            p.append(r)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    ref._p.addprevious(tbl)


def _remove_existing_chapter2(doc: Document) -> None:
    ch3_idx = _find_paragraph(doc, "Chapter 3: System Architecture")
    to_delete = []
    for i in range(ch3_idx - 1, -1, -1):
        text = doc.paragraphs[i].text.strip()
        if text == "Chapter 2: Literature Review":
            to_delete.append(i)
            break
        if any(text.startswith(m) or m in text for m in CH2_MARKERS):
            to_delete.append(i)
        elif text.startswith("2.") and "Architecture" not in text:
            to_delete.append(i)
        elif text.startswith("**BART**") or text.startswith("**PEGASUS**"):
            to_delete.append(i)
        elif text.startswith("| Model |") or text.startswith("| Prior work |"):
            to_delete.append(i)
        elif "Retrieval-Augmented Generation (RAG)" in text:
            to_delete.append(i)
        elif "Prior work addresses individual aspects" in text:
            to_delete.append(i)
        elif "The following table summarises the gap" in text:
            to_delete.append(i)

    ch1_end = _find_paragraph(doc, "Chapter 1: Introduction")
    ch3_idx = _find_paragraph(doc, "Chapter 3: System Architecture")
    for i in range(ch1_end + 1, ch3_idx):
        text = doc.paragraphs[i].text.strip()
        if not text:
            continue
        if any(text.startswith(m) or m in text for m in CH2_MARKERS):
            to_delete.append(i)
        elif re.match(r"^2\.\d", text):
            to_delete.append(i)
        elif "Retrieval-Augmented Generation" in text:
            to_delete.append(i)
        elif "Prior work addresses" in text or "This dissertation fills the gap" in text:
            to_delete.append(i)
        elif text.startswith("**BART**") or text.startswith("| Model |"):
            to_delete.append(i)

    for idx in sorted(set(to_delete), reverse=True):
        if idx < len(doc.paragraphs):
            _delete_paragraph(doc.paragraphs[idx])

    body = doc.element.body
    ch3_idx = _find_paragraph(doc, "Chapter 3: System Architecture")
    anchor_el = doc.paragraphs[ch3_idx]._element
    for tbl in list(body.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")):
        if anchor_el.getprevious() is tbl or tbl in [
            anchor_el.getprevious() for _ in range(5)
        ]:
            parent = tbl.getparent()
            if parent is not None:
                parent.remove(tbl)


def _insert_chapter2(doc: Document) -> None:
    ch3_idx = _find_paragraph(doc, "Chapter 3: System Architecture")
    anchor  = doc.paragraphs[ch3_idx]
    blocks  = _parse_chapter2_blocks()

    for kind, payload in blocks:
        if kind in ("heading1", "heading2", "body"):
            style = "Heading 1" if kind == "heading1" else ("Heading 2" if kind == "heading2" else None)
            _insert_paragraph_before(anchor, payload, style)
        elif kind == "table":
            spacer = _insert_paragraph_before(anchor, "")
            _insert_table_before(spacer, payload)


def _remove_duplicate_contributions(doc: Document) -> None:
    to_delete = []
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t == "1.2 Contributions":
            to_delete.append(i)
            if i + 1 < len(doc.paragraphs) and doc.paragraphs[i + 1].text.strip() == CONTRIBUTIONS:
                to_delete.append(i + 1)
        elif t == CONTRIBUTIONS:
            to_delete.append(i)
    for idx in sorted(set(to_delete), reverse=True):
        _delete_paragraph(doc.paragraphs[idx])


def _ensure_contributions(doc: Document) -> None:
    lit_idx = _find_paragraph(doc, "1.3 Literature Context")
    anchor  = doc.paragraphs[lit_idx]
    contrib = _insert_paragraph_before(anchor, CONTRIBUTIONS)
    _insert_paragraph_before(contrib, "1.2 Contributions", "Heading 2")
    for run in contrib.runs:
        run.italic = True


def _update_chapter1(doc: Document) -> None:
    _replace_list_items(doc, "Objectives", ("Scope of Work",), OBJECTIVES_TOP)

    if not any(TARGETS_NOTE in p.text for p in doc.paragraphs):
        scope_idx = _find_paragraph(doc, "Scope of Work")
        _insert_paragraph_before(doc.paragraphs[scope_idx], TARGETS_NOTE)

    for para in doc.paragraphs:
        t = para.text.strip()
        if t == "1.2 Literature Review and 2026 Context":
            _set_paragraph_text(para, "1.3 Literature Context (see Chapter 2)")
        elif t == "1.2 Stated Objectives (from outline submission)":
            _set_paragraph_text(para, "1.4 Stated Objectives (from outline submission)")
        elif t == "1.3 Objectives Met till Midterm":
            _set_paragraph_text(para, "1.5 Objectives Met till Midterm")

    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if t.startswith("A 2026 multi-track evaluation"):
            _set_paragraph_text(doc.paragraphs[i], LITERATURE_POINTER)
        elif t.startswith("A complementary 2026") or t.startswith("Together, these"):
            _set_paragraph_text(doc.paragraphs[i], "")

    _replace_list_items(
        doc,
        "1.4 Stated Objectives",
        ("1.5 Objectives Met",),
        OBJECTIVES_CH1,
    )


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    if not SRC_PATH.exists():
        raise FileNotFoundError(f"Source report not found: {SRC_PATH}")

    # Start from a fresh copy so we don't mutate the original
    shutil.copy2(str(SRC_PATH), str(OUT_PATH))
    print(f"Copied: {SRC_PATH.name}  ->  {OUT_PATH.name}")

    doc = Document(str(OUT_PATH))

    _renumber_chapters(doc)
    _remove_existing_chapter2(doc)
    _remove_duplicate_contributions(doc)
    _update_chapter1(doc)
    _ensure_contributions(doc)
    _insert_chapter2(doc)

    doc.save(str(OUT_PATH))
    print(f"Saved:  {OUT_PATH}")
    print("Done — MidSem_Report_2024AA05606_ag.docx is ready.")


if __name__ == "__main__":
    main()
