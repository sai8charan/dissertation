"""Apply final styles to the restructured document."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _ensure_child(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el

def _shade_cell(cell, fill):
    tcPr = _ensure_child(cell._tc, "w:tcPr")
    old  = tcPr.find(qn("w:shd"))
    if old is not None:
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def _set_table_borders(table):
    tbl   = table._tbl
    tblPr = _ensure_child(tbl, "w:tblPr")
    old   = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)

doc = Document("documentation/MidSem_Report_2024AA05606_restructured.docx")

# Style definitions (safe access)
try:
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
except KeyError:
    pass

try:
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
except KeyError:
    pass

try:
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
except KeyError:
    pass

try:
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.bold = True
    h3.font.italic = True
    h3.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
except KeyError:
    pass

# Apply directly to paragraphs as well (to override hardcoded run properties and handle ALL CAPS for H1)
for p in doc.paragraphs:
    if p.style is None:
        continue
    
    style_name = p.style.name
    
    if style_name == 'Normal':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = 'Times New Roman'
            
    elif style_name == 'Heading 1':
        p.text = p.text.upper() # ALL CAPS
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    elif style_name == 'Heading 2':
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
            
    elif style_name == 'Heading 3':
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)

# Table formatting
for table in doc.tables:
    _set_table_borders(table)
    # Header row
    for cell in table.rows[0].cells:
        _shade_cell(cell, "E0E0E0") # Light gray
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.name = 'Times New Roman'
                
    # Other rows
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'

# Header formatting
for section in doc.sections:
    header = section.header
    for p in header.paragraphs:
        for r in p.runs:
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80) # Gray
            r.font.name = 'Times New Roman'

out_path = "documentation/MidSem_Report_2024AA05606_styled.docx"
doc.save(out_path)
print(f"Saved styled document to {out_path}")
